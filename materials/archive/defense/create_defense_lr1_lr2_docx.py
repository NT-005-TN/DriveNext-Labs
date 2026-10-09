from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

OUTPUT = "/Users/trukhina12345/Documents/ChatGPT/labs/Защита_ЛР1_ЛР2_DriveNext.docx"

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.06

for style_name, size, before, after in [
    ("Title", 24, 0, 16),
    ("Heading 1", 18, 16, 8),
    ("Heading 2", 14, 12, 6),
    ("Heading 3", 12, 10, 4),
]:
    s = styles[style_name]
    s.font.name = "Arial"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

if "File Path" not in styles:
    st = styles.add_style("File Path", WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = "Courier New"
    st.font.size = Pt(9.5)
    st.font.color.rgb = RGBColor(36, 64, 98)
    st.paragraph_format.left_indent = Inches(0.18)
    st.paragraph_format.space_after = Pt(4)

if "Speech" not in styles:
    st = styles.add_style("Speech", WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = "Arial"
    st.font.size = Pt(10.5)
    st.paragraph_format.left_indent = Inches(0.24)
    st.paragraph_format.right_indent = Inches(0.12)
    st.paragraph_format.space_after = Pt(6)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)

def add_label_paragraph(label, text, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_together = True
    r = p.add_run(label)
    r.bold = True
    p.add_run(text)
    return p

def add_step(title, path, lines, speech, feature=None, connection=None):
    doc.add_heading(title, level=3)
    add_label_paragraph("Открыть файл: ", path, "File Path")
    add_label_paragraph("Показать строки: ", lines)
    add_label_paragraph("Что сказать: ", speech, "Speech")
    if feature:
        add_label_paragraph("Особенность кода: ", feature)
    if connection:
        add_label_paragraph("Связь с другими файлами: ", connection)

def add_bullets(items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)

def add_numbered(items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = header.add_run("DriveNext  сценарий защиты лабораторных работ")
run.font.name = "Arial"
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(90, 90, 90)
add_page_number(section.footer.paragraphs[0])

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Сценарий защиты лабораторных работ 1 и 2")
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p2.add_run("Мобильное приложение DriveNext")
r.bold = True
r.font.size = Pt(16)

doc.add_paragraph()
intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
intro.add_run(
    "Документ предназначен для последовательного показа проекта в Android Studio. "
    "Для каждого этапа указаны папка, файл, точные строки, готовый текст выступления, "
    "особенности Kotlin и Android, а также связь с другими файлами. После разбора кода "
    "приведён порядок демонстрации работоспособности приложения."
)

doc.add_paragraph()
add_label_paragraph("Проект ЛР 1: ", "/Users/trukhina12345/Documents/ChatGPT/labs/DriveNextLab1")
add_label_paragraph("Проект ЛР 2: ", "/Users/trukhina12345/Documents/ChatGPT/labs/DriveNextLab2")

doc.add_heading("Как пользоваться документом", level=1)
add_numbered([
    "В Android Studio включить номера строк через Settings → Editor → General → Appearance → Show line numbers.",
    "Открывать файлы в указанном порядке и выделять только рекомендованный диапазон строк.",
    "Текст после пометки «Что сказать» можно произносить почти дословно.",
    "Папки build, .gradle и .idea не показывать: это автоматически созданные служебные файлы.",
    "После разбора кода перейти к демонстрации приложения на эмуляторе или устройстве.",
])

doc.add_page_break()

# LAB 1
doc.add_heading("Лабораторная работа 1", level=1)
add_label_paragraph("Открыть проект: ", "/Users/trukhina12345/Documents/ChatGPT/labs/DriveNextLab1", "File Path")
add_label_paragraph(
    "Вступление: ",
    "В первой лабораторной работе я реализовала стартовый сценарий приложения DriveNext: "
    "экран загрузки, проверку интернет-соединения, экран отсутствия сети и onboarding из трёх страниц. "
    "Также сохраняется информация о прохождении onboarding, поэтому он показывается только при первом запуске.",
    "Speech",
)

doc.add_heading("Конфигурация проекта", level=2)
add_step(
    "Минимальная версия и View Binding",
    "DriveNextLab1/app/build.gradle.kts",
    "6–17, затем 20–23",
    "Здесь находится конфигурация Android-приложения. compileSdk и targetSdk равны 35, а minSdk равен 31, "
    "то есть минимальная версия — Android 12. В строках 20–23 включён View Binding: Android автоматически "
    "создаёт binding-класс для каждого XML-макета, поэтому обращаться к элементам можно без findViewById.",
    "Название ActivitySplashBinding автоматически образуется от activity_splash.xml.",
)
add_step(
    "Подключённые библиотеки",
    "DriveNextLab1/app/build.gradle.kts",
    "32–42",
    "Здесь подключены Material Components, ViewModel, lifecycleScope, системный SplashScreen, ConstraintLayout и ViewPager2. "
    "Каждая библиотека соответствует отдельной части задания.",
)

doc.add_heading("Манифест и точка входа", level=2)
add_step(
    "Разрешения и стартовая Activity",
    "DriveNextLab1/app/src/main/AndroidManifest.xml",
    "4–5, затем 14–26",
    "INTERNET разрешает сетевые запросы, а ACCESS_NETWORK_STATE позволяет читать состояние сети. "
    "Стартовой является SplashActivity, потому что у неё указаны MAIN и LAUNCHER. Android запускает её при нажатии на иконку приложения.",
    "exported=true разрешает системе запускать компонент через launcher intent-filter.",
)

doc.add_heading("Экран загрузки", level=2)
add_step(
    "XML разметка SplashScreen",
    "DriveNextLab1/app/src/main/res/layout/activity_splash.xml",
    "2–7, 9–30 и 32–44",
    "Корневым контейнером является ConstraintLayout. В верхней части находятся название и слоган, а в центре — иллюстрация. "
    "Тексты вынесены в strings.xml, цвета — в colors.xml. Размер изображения задаётся в процентах от ширины экрана, "
    "поэтому разметка адаптируется под разные устройства.",
    "Значение ширины 0dp в ConstraintLayout означает, что реальный размер определяется ограничениями.",
)

doc.add_heading("Проверка сети", level=2)
add_step(
    "Получение системного ConnectivityManager",
    "DriveNextLab1/app/src/main/java/ru/mtuci/drivenext/data/connectivity/NetworkMonitor.kt",
    "7–9",
    "Класс получает системный ConnectivityManager через applicationContext. Application context используется, чтобы класс не удерживал Activity и не создавал утечку памяти.",
)
add_step(
    "Проверка реального доступа в интернет",
    "DriveNextLab1/app/src/main/java/ru/mtuci/drivenext/data/connectivity/NetworkMonitor.kt",
    "11–16",
    "Сначала получается активная сеть. Если её нет, оператор Элвиса сразу возвращает false. Затем проверяются две возможности: "
    "NET_CAPABILITY_INTERNET означает поддержку интернета, а NET_CAPABILITY_VALIDATED подтверждает, что Android действительно получил доступ во внешнюю сеть.",
    "Оператор ?: является частью null-safety Kotlin и заменяет ручную проверку на null.",
    "NetworkMonitor используется стартовым use case и экраном повторной проверки соединения.",
)

doc.add_heading("Сохранение состояния onboarding", level=2)
add_step(
    "Чтение и запись SharedPreferences",
    "DriveNextLab1/app/src/main/java/ru/mtuci/drivenext/data/preferences/OnboardingPreferences.kt",
    "5–17",
    "Класс открывает приватное хранилище SharedPreferences. Метод isCompleted читает логическое значение, по умолчанию false. "
    "Метод markCompleted записывает true через apply, поэтому сохранение выполняется асинхронно и не блокирует интерфейс.",
    "Имена файла и ключа находятся в companion object и существуют в одном экземпляре на класс.",
    "Значение читает ResolveStartDestinationUseCase, а записывает OnboardingViewModel.",
)

doc.add_heading("Бизнес логика стартовой навигации", level=2)
add_step(
    "Use case выбора экрана",
    "DriveNextLab1/app/src/main/java/ru/mtuci/drivenext/domain/ResolveStartDestinationUseCase.kt",
    "6–17",
    "Use case получает NetworkMonitor и OnboardingPreferences. Выражение when проверяет условия по порядку: без интернета возвращается NO_CONNECTION, "
    "при непройденном onboarding — ONBOARDING, иначе — LOGIN. Возможные результаты описаны enum-классом, поэтому вместо произвольных строк используются типобезопасные значения.",
    "when здесь является выражением и сразу возвращает StartDestination.",
    "Слой data предоставляет факты, domain принимает решение, presentation показывает нужный экран.",
)
add_step(
    "ViewModel стартового экрана",
    "DriveNextLab1/app/src/main/java/ru/mtuci/drivenext/presentation/splash/SplashViewModel.kt",
    "7–8",
    "ViewModel не проверяет сеть самостоятельно, а делегирует решение use case. Это отделяет интерфейс от бизнес-логики.",
)

doc.add_heading("Запуск приложения", level=2)
add_step(
    "Создание ViewModel с зависимостями",
    "DriveNextLab1/app/src/main/java/ru/mtuci/drivenext/presentation/splash/SplashActivity.kt",
    "22–33",
    "ViewModel создаётся через делегат by viewModels. Поскольку у неё есть параметр конструктора, применяется ViewModelProvider.Factory. "
    "Фабрика создаёт use case и передаёт ему объекты слоя данных.",
    "by viewModels создаёт ViewModel лениво и сохраняет её при изменении конфигурации.",
)
add_step(
    "SplashScreen, coroutine и переход",
    "DriveNextLab1/app/src/main/java/ru/mtuci/drivenext/presentation/splash/SplashActivity.kt",
    "35–53",
    "Сначала устанавливается системный SplashScreen, затем подключается XML через View Binding. В lifecycleScope запускается coroutine, "
    "которая ждёт 2500 миллисекунд. После задержки ViewModel возвращает следующий экран. Intent открывает выбранную Activity, а finish удаляет splash из стека.",
    "lifecycleScope отменяет coroutine, если Activity уничтожается. Время вынесено в константу SPLASH_DURATION_MS.",
)

doc.add_heading("Экран отсутствия сети", level=2)
add_step(
    "Разметка сообщения и кнопки",
    "DriveNextLab1/app/src/main/res/layout/activity_no_connection.xml",
    "20–43 и 46–58",
    "Здесь находятся иконка, два текста и кнопка Повторить попытку. Идентификатор retryButton превращается в свойство binding.retryButton.",
)
add_step(
    "Повторная проверка соединения",
    "DriveNextLab1/app/src/main/java/ru/mtuci/drivenext/presentation/connection/NoConnectionActivity.kt",
    "27–43",
    "Activity подключает разметку через ActivityNoConnectionBinding. По нажатию кнопки ViewModel повторно проверяет интернет. "
    "Если сеть появилась, приложение читает состояние onboarding и выбирает экран. Если сети нет, показывается Snackbar.",
)

doc.add_heading("Onboarding", level=2)
add_step(
    "Список страниц",
    "DriveNextLab1/app/src/main/java/ru/mtuci/drivenext/presentation/onboarding/OnboardingViewModel.kt",
    "7–14",
    "ViewModel содержит три объекта OnboardingPage: аренда, безопасность и предложения. Каждый объект хранит ресурсы изображения, заголовка и описания. "
    "Метод complete сохраняет прохождение через OnboardingPreferences.",
)
add_step(
    "Адаптер ViewPager2",
    "DriveNextLab1/app/src/main/java/ru/mtuci/drivenext/presentation/onboarding/OnboardingAdapter.kt",
    "6–9",
    "FragmentStateAdapter сообщает ViewPager2 количество страниц и создаёт фрагмент для выбранной позиции.",
    "Адаптер связывает список из ViewModel с визуальными фрагментами.",
)
add_step(
    "Безопасная работа с binding во Fragment",
    "DriveNextLab1/app/src/main/java/ru/mtuci/drivenext/presentation/onboarding/OnboardingPageFragment.kt",
    "10–27",
    "Binding хранится в nullable-поле, потому что жизненный цикл View короче жизненного цикла Fragment. В onViewCreated данные из arguments устанавливаются в интерфейс, "
    "а в onDestroyView ссылка очищается, чтобы не удерживать уничтоженную View.",
)
add_step(
    "Передача данных через Bundle",
    "DriveNextLab1/app/src/main/java/ru/mtuci/drivenext/presentation/onboarding/OnboardingPageFragment.kt",
    "30–40",
    "Фабричный метод newInstance создаёт Fragment и передаёт идентификаторы ресурсов через Bundle. Поэтому все слайды используют один Fragment и одну XML-разметку.",
)
add_step(
    "Управление перелистыванием",
    "DriveNextLab1/app/src/main/java/ru/mtuci/drivenext/presentation/onboarding/OnboardingActivity.kt",
    "29–45",
    "Адаптер подключается к ViewPager2. Callback реагирует на смену страницы. Кнопка Пропустить сразу завершает onboarding, а кнопка Далее увеличивает позицию. "
    "На последней странице вызывается завершение.",
)
add_step(
    "Индикаторы и завершение",
    "DriveNextLab1/app/src/main/java/ru/mtuci/drivenext/presentation/onboarding/OnboardingActivity.kt",
    "54–75",
    "Если открыта последняя страница, текст кнопки меняется на Поехали. Индикаторы создаются динамически: активный — вытянутый фиолетовый, остальные — маленькие серые. "
    "При завершении ViewModel сохраняет состояние, запускается LoginActivity, а onboarding закрывается.",
    "Функция dp переводит dp в реальные пиксели с учётом плотности экрана.",
)

doc.add_heading("Демонстрация лабораторной работы 1", level=2)
add_numbered([
    "Очистить данные приложения, чтобы onboarding считался непройденным.",
    "Запустить приложение и показать стартовый экран в течение 2,5 секунды.",
    "Перелистнуть onboarding жестом и кнопкой Далее.",
    "На последнем слайде показать изменение текста кнопки на Поехали.",
    "Завершить onboarding и перезапустить приложение: onboarding больше не должен появиться.",
    "Отключить интернет и перезапустить приложение: должен открыться экран отсутствия сети.",
    "Включить интернет и нажать Повторить попытку: приложение должно продолжить навигацию.",
])
add_label_paragraph(
    "Финальная фраза: ",
    "Во время демонстрации видно, что слой данных проверяет сеть и хранит состояние, domain-слой выбирает сценарий, "
    "а Activity и Fragment отвечают только за отображение и пользовательские действия.",
    "Speech",
)

doc.add_page_break()

# LAB 2
doc.add_heading("Лабораторная работа 2", level=1)
add_label_paragraph("Открыть проект: ", "/Users/trukhina12345/Documents/ChatGPT/labs/DriveNextLab2", "File Path")
add_label_paragraph(
    "Вступление: ",
    "Во второй лабораторной работе я реализовала выбор между входом и регистрацией, форму авторизации и регистрацию из трёх шагов. "
    "В работе используются проверка email, показ и скрытие пароля, MaterialDatePicker, Activity Result API, Uri и сохранение промежуточных данных.",
    "Speech",
)

doc.add_heading("Точка входа", level=2)
add_step(
    "Регистрация экранов",
    "DriveNextLab2/app/src/main/AndroidManifest.xml",
    "14–30",
    "В манифесте зарегистрированы главный экран, успешная регистрация, три шага регистрации и вход. Стартовой во второй работе является AuthChoiceActivity, "
    "потому что у неё указаны MAIN и LAUNCHER.",
)

doc.add_heading("Выбор входа или регистрации", level=2)
add_step(
    "Разметка первого экрана",
    "DriveNextLab2/app/src/main/res/layout/activity_auth_choice.xml",
    "3–7",
    "В XML находятся название, слоган, иллюстрация и две кнопки. ID signInButton и registerButton используются View Binding для создания одноимённых свойств.",
)
add_step(
    "Обработка двух переходов",
    "DriveNextLab2/app/src/main/java/ru/mtuci/drivenext/presentation/auth/AuthChoiceActivity.kt",
    "9–15",
    "В onCreate создаётся binding. Первая кнопка открывает LoginActivity, вторая — RegisterStep1Activity. Переходы выполняются через Intent.",
)

doc.add_heading("Авторизация", level=2)
add_step(
    "Поля email и пароля",
    "DriveNextLab2/app/src/main/res/layout/activity_login.xml",
    "6–11",
    "Email использует inputType textEmailAddress. Пароль использует textPassword, а TextInputLayout автоматически добавляет иконку глаза через password_toggle. "
    "Кнопка входа изначально отключена атрибутом enabled=false.",
)
add_step(
    "Переиспользуемый TextWatcher",
    "DriveNextLab2/app/src/main/java/ru/mtuci/drivenext/presentation/common/TextWatchers.kt",
    "7–12",
    "Это extension-функция для EditText. Интерфейс TextWatcher требует три метода, но приложению нужен только afterTextChanged. "
    "Параметр block имеет тип функции без аргументов и результата и вызывается после изменения текста.",
    "Extension-функция вызывается как обычный метод поля: emailInput.afterTextChanged.",
)
add_step(
    "Включение кнопки и проверка email",
    "DriveNextLab2/app/src/main/java/ru/mtuci/drivenext/presentation/auth/LoginActivity.kt",
    "14–30",
    "Локальная функция updateState включает кнопку только при заполнении обоих полей. Одна и та же функция передаётся двум TextWatcher через ссылку ::updateState. "
    "Email проверяется стандартным Patterns.EMAIL_ADDRESS. При успешной демонстрационной авторизации в SharedPreferences сохраняется demo_token, "
    "а флаги NEW_TASK и CLEAR_TASK очищают стек входа.",
    "Серверная авторизация и Google OAuth требуют внешних ключей; в лабораторной реализован клиентский сценарий и валидация.",
)

doc.add_heading("Черновик многошаговой регистрации", level=2)
add_step(
    "Singleton с введёнными данными",
    "DriveNextLab2/app/src/main/java/ru/mtuci/drivenext/data/registration/RegistrationDraft.kt",
    "3–15",
    "Ключевое слово object создаёт singleton — один экземпляр на приложение. В нём временно хранятся данные всех шагов. Метод clear очищает черновик после завершения регистрации.",
    "Это временное хранилище состояния, а не постоянная база данных.",
)

doc.add_heading("Первый шаг регистрации", level=2)
add_step(
    "XML первого шага",
    "DriveNextLab2/app/src/main/res/layout/activity_register_step1.xml",
    "6–11",
    "Первый шаг содержит email, два поля пароля, обязательный CheckBox и кнопку Далее. Для обоих паролей включён password_toggle.",
)
add_step(
    "Восстановление и валидация",
    "DriveNextLab2/app/src/main/java/ru/mtuci/drivenext/presentation/registration/RegisterStep1Activity.kt",
    "15–32",
    "Сначала значения восстанавливаются из RegistrationDraft. После нажатия Далее email очищается от пробелов через trim. Выражение when проверяет email, "
    "заполненность паролей, их совпадение и CheckBox. Если ошибок нет, данные сохраняются в черновик и открывается второй шаг.",
)

doc.add_heading("Общий MaterialDatePicker", level=2)
add_step(
    "Создание календаря",
    "DriveNextLab2/app/src/main/java/ru/mtuci/drivenext/presentation/common/DatePicker.kt",
    "10–16",
    "Extension-функция создаёт MaterialDatePicker через Builder. После подтверждения календарь возвращает время в миллисекундах UTC. "
    "SimpleDateFormat преобразует его в строку переданного формата, после чего вызывается callback onSelected.",
    "Один код используется для даты рождения и даты выдачи; отличается только параметр pattern.",
)

doc.add_heading("Второй шаг регистрации", level=2)
add_step(
    "XML персональных данных",
    "DriveNextLab2/app/src/main/res/layout/activity_register_step2.xml",
    "6–13",
    "Здесь расположены фамилия, имя, необязательное отчество, дата рождения и RadioGroup выбора пола. Поле даты имеет focusable=false, "
    "поэтому вместо клавиатуры открывается MaterialDatePicker.",
)
add_step(
    "Восстановление данных, календарь и проверка",
    "DriveNextLab2/app/src/main/java/ru/mtuci/drivenext/presentation/registration/RegisterStep2Activity.kt",
    "15–36",
    "Activity восстанавливает значения из черновика и отмечает ранее выбранный RadioButton. По нажатию на дату вызывается showDatePicker с форматом MM/dd/yyyy. "
    "Пол определяется по checkedRadioButtonId. После проверки данные записываются в RegistrationDraft и открывается третий шаг.",
)

doc.add_heading("Третий шаг регистрации", level=2)
add_step(
    "XML документов",
    "DriveNextLab2/app/src/main/res/layout/activity_register_step3.xml",
    "6–13",
    "На экране находятся фото профиля, номер удостоверения, дата выдачи и две обязательные загрузки документов. Для номера установлено maxLength=10.",
)
add_step(
    "Регистрация системного пикера",
    "DriveNextLab2/app/src/main/java/ru/mtuci/drivenext/presentation/registration/RegisterStep3Activity.kt",
    "14–26",
    "Activity хранит текущую цель выбора и два nullable Uri обязательных документов. Контракт PickVisualMedia регистрируется заранее. "
    "Если пользователь отменяет выбор, возвращается null. Иначе when направляет Uri в фото профиля, удостоверения или паспорта.",
    "Uri является безопасной ссылкой на выбранный файл; приложению не нужен полный доступ ко всей галерее.",
)
add_step(
    "Запуск пикера и валидация документов",
    "DriveNextLab2/app/src/main/java/ru/mtuci/drivenext/presentation/registration/RegisterStep3Activity.kt",
    "32–52",
    "Три кнопки вызывают одну функцию pick, но передают разные значения enum PhotoTarget. При продолжении проверяются длина номера, наличие даты и двух обязательных фото. "
    "Фото профиля не проверяется, потому что по заданию оно необязательное.",
    "Enum ограничивает назначение изображения тремя допустимыми вариантами.",
)

doc.add_heading("Успешная регистрация", level=2)
add_step(
    "Разметка результата",
    "DriveNextLab2/app/src/main/res/layout/activity_success.xml",
    "3–7",
    "Экран содержит заголовок, векторную иконку, сообщение и кнопку. ConstraintLayout сохраняет взаимное расположение элементов на разных экранах.",
)
add_step(
    "Очистка данных и стека",
    "DriveNextLab2/app/src/main/java/ru/mtuci/drivenext/presentation/registration/SuccessActivity.kt",
    "10–18",
    "По нажатию кнопки RegistrationDraft очищается. Главный экран запускается с NEW_TASK и CLEAR_TASK, поэтому пользователь не может вернуться к персональным данным кнопкой Назад.",
)

doc.add_heading("Демонстрация лабораторной работы 2", level=2)
add_numbered([
    "На стартовом экране выбрать Зарегистрироваться.",
    "Ввести некорректный email и показать сообщение проверки.",
    "Ввести разные пароли и показать ошибку Пароли не совпадают.",
    "Не ставить CheckBox и показать требование согласия с условиями.",
    "Исправить первый шаг и перейти ко второму.",
    "Оставить обязательное поле пустым и показать Snackbar.",
    "Выбрать дату рождения через MaterialDatePicker и выбрать пол.",
    "На третьем шаге ввести номер короче 10 символов и показать ошибку.",
    "Ввести 1234567890, выбрать дату выдачи и две фотографии документов.",
    "Завершить регистрацию и показать экран успеха.",
    "Нажать Далее и показать, что стек регистрации очищен.",
    "Отдельно открыть экран входа, показать неактивную кнопку при пустых полях и её включение после ввода email и пароля.",
])

doc.add_heading("Итоговый ответ преподавателю", level=2)
add_label_paragraph(
    "Что сказать: ",
    "В первой лабораторной работе я реализовала SplashScreen, проверку сети, ViewPager2 и сохранение прохождения onboarding. "
    "Во второй работе реализовала авторизацию и многошаговую регистрацию с валидацией, MaterialDatePicker, Activity Result API и Uri. "
    "XML отвечает за разметку, Activity и Fragment — за пользовательские действия, ViewModel — за связь с логикой, а классы data и domain — за данные и принятие решений.",
    "Speech",
)

# Format paragraphs and tables consistently.
for p in doc.paragraphs:
    p.paragraph_format.widow_control = True
    for run in p.runs:
        if not run.font.name:
            run.font.name = "Arial"

doc.core_properties.title = "Сценарий защиты лабораторных работ 1 и 2 DriveNext"
doc.core_properties.subject = "Показ кода и демонстрация Android приложения"
doc.core_properties.author = "Студент"
doc.save(OUTPUT)
print(OUTPUT)
